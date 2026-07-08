from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "IWA_Technical_Documentation.md"
CFG_TS = ROOT / "src" / "webparts" / "iwa" / "components" / "data" / "cfg.ts"
OUTPUT_DOCX = ROOT / "docs" / "IWA_Technical_Documentation.docx"


LIST_NAME_MAP = {
    "Authorizations": "IWAAgreements",
    "Counters": "IWACounters",
    "Exports": "IWAExports",
    "LaborLine": "IWALaborLine",
    "Mods": "IWAMods",
    "Resources": "IWAResources",
    "TravelODC": "IWATravelODC",
    "Users": "IWAAppUsers",
    "WorkflowRuns": "IWAWorkflowRuns",
    "WorkflowActions": "IWAWorkflowActions",
}


def iter_code_aware(text: str, start: int = 0):
    i = start
    quote = None
    line_comment = False
    block_comment = False
    while i < len(text):
        ch = text[i]
        nxt = text[i + 1] if i + 1 < len(text) else ""
        if line_comment:
            if ch == "\n":
                line_comment = False
            i += 1
            continue
        if block_comment:
            if ch == "*" and nxt == "/":
                block_comment = False
                i += 2
            else:
                i += 1
            continue
        if quote:
            if ch == "\\":
                i += 2
                continue
            if ch == quote:
                quote = None
            i += 1
            continue
        if ch == "/" and nxt == "/":
            line_comment = True
            i += 2
            continue
        if ch == "/" and nxt == "*":
            block_comment = True
            i += 2
            continue
        if ch in ("'", '"', "`"):
            quote = ch
            i += 1
            continue
        yield i, ch
        i += 1


def find_matching(text: str, open_index: int) -> int:
    pairs = {"{": "}", "[": "]", "(": ")"}
    opener = text[open_index]
    closer = pairs[opener]
    depth = 0
    for i, ch in iter_code_aware(text, open_index):
        if ch == opener:
            depth += 1
        elif ch == closer:
            depth -= 1
            if depth == 0:
                return i
    raise ValueError(f"No matching {closer} for index {open_index}")


def find_backward_object_start(text: str, pos: int) -> int:
    for i in range(pos, -1, -1):
        if text[i] == "{":
            return i
    raise ValueError("Could not find object start.")


def find_array_after(text: str, token: str) -> str:
    token_pos = text.find(token)
    if token_pos < 0:
        return ""
    open_index = text.find("[", token_pos)
    if open_index < 0:
        return ""
    close_index = find_matching(text, open_index)
    return text[open_index + 1:close_index]


def split_top_level_objects(array_body: str) -> list[str]:
    objects = []
    depth = 0
    start = None
    for i, ch in iter_code_aware(array_body):
        if ch == "{":
            if depth == 0:
                start = i
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and start is not None:
                objects.append(array_body[start:i + 1])
                start = None
    return objects


def clean_value(value: str) -> str:
    value = value.strip().rstrip(",")
    value = re.sub(r"\s+as\s+Helper\.[A-Za-z0-9_]+", "", value).strip()
    if value.startswith('"') and value.endswith('"'):
        return value[1:-1]
    if value.startswith("'") and value.endswith("'"):
        return value[1:-1]
    if value.startswith("[") and value.endswith("]"):
        return " ".join(re.findall(r'["\']([^"\']+)["\']', value))
    if "Helper.SPCfgFieldType." in value:
        return value.split("Helper.SPCfgFieldType.", 1)[1].split()[0].strip(",")
    if "SPTypes.DateFormat." in value:
        return value.split("SPTypes.DateFormat.", 1)[1].split()[0].strip(",")
    if "SPTypes.FieldNoteType." in value:
        return value.split("SPTypes.FieldNoteType.", 1)[1].split()[0].strip(",")
    if value.startswith("Strings.Sites.main.lists."):
        key = value.split(".")[-1].strip()
        return LIST_NAME_MAP.get(key, key)
    return value


def extract_prop(block: str, prop: str) -> str:
    match = re.search(rf"\b{re.escape(prop)}\s*:\s*(.+?)(?:,|\n\s*}})", block, re.S)
    if not match:
        return ""
    return clean_value(match.group(1))


def parse_list_schema() -> list[dict]:
    text = CFG_TS.read_text(encoding="utf-8")
    schemas = []
    pos = 0
    while True:
        info_pos = text.find("ListInformation:", pos)
        if info_pos < 0:
            break
        obj_start = find_backward_object_start(text, info_pos)
        obj_end = find_matching(text, obj_start)
        block = text[obj_start:obj_end + 1]
        pos = obj_end + 1

        title_expr = extract_prop(block, "Title")
        title = title_expr if title_expr else "Unknown List"
        description = extract_prop(block, "Description")
        base_template = extract_prop(block, "BaseTemplate").replace("SPTypes.ListTemplateType.", "")
        title_display = extract_prop(block, "TitleFieldDisplayName")
        if not title_display:
            title_display = "Title"

        fields = [
            {
                "internal_name": "Title",
                "display_name": title_display,
                "type": "Text",
                "required": "Yes" if extract_prop(block, "TitleFieldRequired") != "false" else "No",
                "indexed": "",
                "details": "Built-in SharePoint title field.",
            }
        ]

        fields_body = find_array_after(block, "CustomFields")
        for field_block in split_top_level_objects(fields_body):
            choices = " ".join(re.findall(r'choices\s*:\s*\[(.*?)\]', field_block, re.S))
            choice_values = ", ".join(re.findall(r'["\']([^"\']+)["\']', choices))
            list_name = extract_prop(field_block, "listName")
            show_field = extract_prop(field_block, "showField")
            decimals = extract_prop(field_block, "decimals")
            date_format = extract_prop(field_block, "format")
            note_type = extract_prop(field_block, "noteType")
            default = extract_prop(field_block, "defaultValue")
            multi = extract_prop(field_block, "multi")
            details = []
            if choice_values:
                details.append(f"Choices: {choice_values}")
            if list_name:
                details.append(f"Lookup list: {list_name}")
            if show_field:
                details.append(f"Show field: {show_field}")
            if decimals:
                details.append(f"Decimals: {decimals}")
            if date_format:
                details.append(f"Format: {date_format}")
            if note_type:
                details.append(f"Note type: {note_type}")
            if default:
                details.append(f"Default: {default}")
            if multi == "true":
                details.append("Multi-value")

            fields.append(
                {
                    "internal_name": extract_prop(field_block, "name"),
                    "display_name": extract_prop(field_block, "title"),
                    "type": extract_prop(field_block, "type"),
                    "required": "Yes" if extract_prop(field_block, "required") == "true" else "No",
                    "indexed": "Yes" if extract_prop(field_block, "indexed") == "true" else "No",
                    "details": "; ".join(details),
                }
            )

        schemas.append(
            {
                "title": title,
                "description": description,
                "base_template": base_template,
                "fields": fields,
            }
        )
    return schemas


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = width
                tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(width))
                tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.color.rgb = RGBColor.from_string("0B2545")
    title.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.text = "Koniag Intercompany Work Authorization Technical Documentation"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string("666666")
    add_page_number(section.footer.add_paragraph())


def add_table(doc: Document, rows: list[list[str]], widths: list[int], header: bool = True) -> None:
    if not rows:
        return
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_cell_margins(table)
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        if r_idx == 0 and header:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(value or "")
            run.font.size = Pt(8.5)
            if r_idx == 0 and header:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("0B2545")
                set_cell_shading(cell, "E8EEF5")
    set_table_width(table, widths)
    doc.add_paragraph()


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if re.match(r"^\|\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?$", line):
            i += 1
            continue
        cells = [cell.strip().replace("`", "") for cell in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_markdown_content(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    i = 0
    in_code = False
    code_buffer = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style=None)
                run = p.add_run("\n".join(code_buffer))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_markdown_table(lines, i)
            if rows:
                col_count = len(rows[0])
                widths = {
                    2: [2700, 6660],
                    3: [1800, 2100, 5460],
                    4: [1500, 1700, 2100, 4060],
                }.get(col_count, [max(1000, 9360 // col_count)] * col_count)
                add_table(doc, rows, widths)
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline_runs(p, stripped[2:])
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=3)
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)
        i += 1


def add_schema_appendix(doc: Document, schemas: list[dict]) -> None:
    doc.add_page_break()
    doc.add_heading("Appendix A: SharePoint List Fields", level=1)
    doc.add_paragraph(
        "The following tables summarize the installed SharePoint lists and custom fields defined in "
        "src/webparts/iwa/components/data/cfg.ts. The built-in Title field is included because it is "
        "renamed or repurposed on several lists."
    )
    summary_rows = [["List", "Base Template", "Field Count", "Description"]]
    for schema in schemas:
        summary_rows.append([
            schema["title"],
            schema["base_template"],
            str(len(schema["fields"])),
            schema["description"],
        ])
    add_table(doc, summary_rows, [1900, 1600, 1100, 4760])

    for schema in schemas:
        doc.add_heading(schema["title"], level=2)
        if schema["description"]:
            doc.add_paragraph(schema["description"])
        rows = [["Internal Name", "Display Name", "Type", "Req.", "Idx.", "Details"]]
        for field in schema["fields"]:
            rows.append([
                field["internal_name"],
                field["display_name"],
                field["type"],
                field["required"],
                field["indexed"],
                field["details"],
            ])
        add_table(doc, rows, [1500, 1700, 1300, 650, 650, 3560])


def main() -> None:
    doc = Document()
    style_document(doc)
    markdown = SOURCE_MD.read_text(encoding="utf-8")
    add_markdown_content(doc, markdown)
    add_schema_appendix(doc, parse_list_schema())
    doc.core_properties.title = "Koniag Intercompany Work Authorization Technical Documentation"
    doc.core_properties.subject = "Technical documentation and SharePoint list field appendix"
    doc.core_properties.author = "Koniag / Codex"
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
